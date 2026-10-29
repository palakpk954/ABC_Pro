"""
Multi-Format Document Parser Service.

Extracts text and metadata from heterogeneous document formats:
- PDF (via pypdf)
- Word DOCX (via python-docx)
- Structured / Semi-structured (JSON, CSV)
- Plain Text / Markdown
"""

import io
import json
import logging
from typing import Optional, Tuple

logger = logging.getLogger(__name__)

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None


class DocumentParser:
    """Parses text from various file formats with robust error handling and fallback."""

    @classmethod
    def parse_file(
        cls,
        filename: str,
        content: bytes,
        mime_type: Optional[str] = None
    ) -> Tuple[str, str]:
        """
        Parse raw bytes into extracted clean text.
        Returns:
            Tuple of (extracted_text, detected_file_type)
        """
        filename_lower = filename.lower()

        if filename_lower.endswith(".pdf"):
            return cls.parse_pdf(content), "pdf"
        elif filename_lower.endswith(".docx") or filename_lower.endswith(".doc"):
            return cls.parse_docx(content), "docx"
        elif filename_lower.endswith(".json"):
            return cls.parse_json(content), "json"
        elif filename_lower.endswith(".csv"):
            return cls.parse_csv_or_text(content), "csv"
        elif filename_lower.endswith((".txt", ".md", ".markdown")):
            return cls.parse_csv_or_text(content), "text"
        else:
            # Fallback to UTF-8 decoding
            return cls.parse_csv_or_text(content), "unknown"

    @classmethod
    def parse_pdf(cls, content: bytes) -> str:
        """Extract text from PDF using pypdf."""
        if not pypdf:
            logger.warning("pypdf is not available, attempting raw text decode fallback")
            return content.decode("utf-8", errors="replace")

        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            extracted_pages = []
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    extracted_pages.append(f"--- Page {idx + 1} ---\n{page_text}")
            return "\n\n".join(extracted_pages).strip()
        except Exception as e:
            logger.error(f"Failed to parse PDF: {e}")
            return content.decode("utf-8", errors="replace")

    @classmethod
    def parse_docx(cls, content: bytes) -> str:
        """Extract text from Word DOCX using python-docx."""
        if not docx:
            logger.warning("python-docx is not available, attempting raw text decode fallback")
            return content.decode("utf-8", errors="replace")

        try:
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract table text
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            return "\n".join(paragraphs).strip()
        except Exception as e:
            logger.error(f"Failed to parse DOCX: {e}")
            return content.decode("utf-8", errors="replace")

    @classmethod
    def parse_json(cls, content: bytes) -> str:
        """Extract readable representation from JSON."""
        try:
            text = content.decode("utf-8", errors="replace")
            parsed = json.loads(text)
            return json.dumps(parsed, indent=2)
        except Exception:
            return content.decode("utf-8", errors="replace")

    @classmethod
    def parse_csv_or_text(cls, content: bytes) -> str:
        """Safely decode plain text or CSV."""
        return content.decode("utf-8", errors="replace").strip()

    @classmethod
    def infer_document_category(cls, filename: str, content_text: str) -> str:
        """
        Infer whether document is a resume, performance_report, feedback_360,
        activity_log, or availability schedule.
        """
        lower_name = filename.lower()
        lower_text = content_text.lower()[:3000]

        # Name-based heuristics
        if any(w in lower_name for w in ["resume", "cv", "curriculum", "bio"]):
            return "resume"
        if any(w in lower_name for w in ["performance", "review", "appraisal", "annual", "quarterly", "q1", "q2", "q3", "q4"]):
            return "performance_report"
        if any(w in lower_name for w in ["feedback", "360", "peer", "survey", "colleague"]):
            return "feedback_360"
        if any(w in lower_name for w in ["activity", "monthly", "bimonthly", "timesheet", "sprint", "log", "summary"]):
            return "activity_log"
        if any(w in lower_name for w in ["availability", "leave", "pto", "vacation", "calendar", "schedule"]):
            return "availability"

        # Content-based heuristics
        if "education" in lower_text and ("experience" in lower_text or "employment" in lower_text):
            return "resume"
        if any(phrase in lower_text for phrase in ["strengths", "areas of improvement", "performance rating", "key accomplishments", "goal evaluation"]):
            return "performance_report"
        if any(phrase in lower_text for phrase in ["peer feedback", "collaborator feedback", "leadership feedback", "how does this person work"]):
            return "feedback_360"
        if any(phrase in lower_text for phrase in ["tasks completed", "sprint deliveries", "hours logged", "milestones reached"]):
            return "activity_log"
        if any(phrase in lower_text for phrase in ["out of office", "pto balance", "planned leave", "working hours"]):
            return "availability"

        return "resume"  # Default assumption
